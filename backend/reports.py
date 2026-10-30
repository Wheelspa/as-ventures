import io
from datetime import datetime

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer

from core import db, scope, can, inr, today_ist, now_iso, IST, LEDGER_TYPES, notify
from finance import balances, enrich_settlement, settlement_status, compute_pricing

REPORTS = {
    "investment_statement": "Investor-wise Investment Statement",
    "principal_recovery": "Order-wise Principal Recovery Report",
    "investor_margin": "Investor Margin Earned & Paid Report",
    "wheelspa_margin": "Wheelspa Margin Report",
    "inventory": "PPF Inventory & Stock Movement Report",
    "purchases": "Purchase & Supplier Payment Report",
    "sales_collections": "Customer Sales & Collection Report",
    "outstanding_principal": "Outstanding Principal Report",
    "pending_settlements": "Pending Settlement Report",
    "monthly_reconciliation": "Monthly Investor Reconciliation Statement",
    "gst": "GST & Tax-Supporting Transaction Report",
    "audit": "Audit Trail & Transaction Modification Report",
}
SOLD = ["invoiced", "partially_paid", "paid"]


def rng(f, t):
    q = {}
    if f:
        q["$gte"] = f
    if t:
        q["$lte"] = t + "~"
    return q


def base_scope(m, investor_id=None):
    q = scope(m)
    if investor_id and can(m, "staff"):
        q = {"investor_ids": investor_id}
    return q


async def _all(coll, q, sort=None):
    if coll in ("rolls", "investors", "purchases", "investments", "collections"):
        q = {**q, "voided": {"$ne": True}}
    cur = db[coll].find(q, {"_id": 0})
    if sort:
        cur = cur.sort(sort)
    return await cur.to_list(5000)


async def ledger_sums(investor_id, date_q=None):
    match = {"investor_id": investor_id}
    if date_q:
        match["date"] = date_q
    out = {}
    async for r in db.ledger.aggregate([{"$match": match}, {"$group": {"_id": "$type", "s": {"$sum": "$amount"}, "n": {"$sum": 1}}}]):
        out[r["_id"]] = r["s"]
    return out


async def monthly_reco(investor_id, month):
    inv = await db.investors.find_one({"id": investor_id}, {"_id": 0})
    if not inv:
        raise HTTPException(404, "Investor not found")
    start = f"{month}-01"
    y, mo = int(month[:4]), int(month[5:7])
    nxt = f"{y + (mo == 12)}-{(mo % 12) + 1:02d}-01"
    o = await ledger_sums(investor_id, {"$lt": start})
    p = await ledger_sums(investor_id, {"$gte": start, "$lt": nxt})
    g = lambda d, k: d.get(k, 0)
    rows = []

    def row(label, open_, add, ded):
        rows.append({"particular": label, "opening": open_, "additions": add, "deductions": ded, "closing": open_ + add - ded})

    row("Principal outstanding (capital received − repaid)", g(o, "CAPITAL_RECEIVED") - g(o, "PRINCIPAL_REPAYMENT"), g(p, "CAPITAL_RECEIVED"), g(p, "PRINCIPAL_REPAYMENT"))
    row("Available funds for procurement", g(o, "CAPITAL_RECEIVED") - g(o, "FUNDS_ALLOCATED"), g(p, "CAPITAL_RECEIVED"), g(p, "FUNDS_ALLOCATED"))
    row("Capital deployed in inventory", g(o, "FUNDS_ALLOCATED") - g(o, "PRINCIPAL_RECOVERED"), g(p, "FUNDS_ALLOCATED"), g(p, "PRINCIPAL_RECOVERED"))
    row("Principal recovered, pending transfer", g(o, "PRINCIPAL_RECOVERED") - g(o, "PRINCIPAL_REPAYMENT"), g(p, "PRINCIPAL_RECOVERED"), g(p, "PRINCIPAL_REPAYMENT"))
    row("Investor margin payable", g(o, "INVESTOR_MARGIN_EARNED") - g(o, "INVESTOR_MARGIN_PAID"), g(p, "INVESTOR_MARGIN_EARNED"), g(p, "INVESTOR_MARGIN_PAID"))
    sales_n = await db.ledger.count_documents({"type": "SALE_INVOICED", "investor_ids": investor_id, "date": {"$gte": start, "$lt": nxt}})
    summary = {"Investor": f"{inv['name']} ({inv['code']})", "Month": month, "Receipts": inr(g(p, "CAPITAL_RECEIVED")),
               "Procurement allocation": inr(g(p, "FUNDS_ALLOCATED")), "Rolls invoiced": sales_n,
               "Principal repaid": inr(g(p, "PRINCIPAL_REPAYMENT")), "Margin earned": inr(g(p, "INVESTOR_MARGIN_EARNED")),
               "Margin paid": inr(g(p, "INVESTOR_MARGIN_PAID"))}
    cols = [["particular", "Particulars", "text"], ["opening", "Opening", "money"], ["additions", "Additions", "money"],
            ["deductions", "Deductions", "money"], ["closing", "Closing", "money"]]
    return cols, rows, summary


async def build_report(key, m, f=None, t=None, investor_id=None, month=None):
    if key not in REPORTS:
        raise HTTPException(404, "Unknown report")
    q = base_scope(m, investor_id)
    summary = {}
    if key == "investment_statement":
        rows = await _all("investments", {**q, **({"date": rng(f, t)} if f or t else {})}, [("date", 1)])
        cols = [["date", "Date", "text"], ["ref_no", "Ref", "text"], ["investor_name", "Investor", "text"], ["amount", "Amount", "money"],
                ["bank_ref", "Bank ref", "text"], ["terms_note", "Terms", "text"]]
        summary = {"Total received": inr(sum(r["amount"] for r in rows))}
    elif key in ("principal_recovery", "investor_margin", "pending_settlements"):
        sq = {**q, "active": True}
        if f or t:
            sq["last_collection_date"] = rng(f, t)
        rows = [await enrich_settlement(s) for s in await _all("settlements", sq, [("order_no", 1)])]
        if key == "pending_settlements":
            rows = [r for r in rows if r["status"] != "settled"]
        for r in rows:
            r["principal_pending"] = r["principal_recoverable"] - r["principal_paid"]
            r["margin_payable"] = r["margin_recoverable"] - r["margin_paid"]
        if key == "investor_margin":
            cols = [["order_no", "Order", "text"], ["roll_serial", "Roll", "text"], ["investor_name", "Investor", "text"],
                    ["margin_due", "Margin due", "money"], ["margin_recoverable", "Earned (collected)", "money"],
                    ["margin_paid", "Paid (verified)", "money"], ["margin_payable", "Payable", "money"], ["status", "Status", "text"]]
            summary = {"Earned": inr(sum(r["margin_recoverable"] for r in rows)), "Paid": inr(sum(r["margin_paid"] for r in rows))}
        else:
            cols = [["order_no", "Order", "text"], ["roll_serial", "Roll", "text"], ["investor_name", "Investor", "text"],
                    ["principal_due", "Principal due", "money"], ["principal_recoverable", "Recoverable (collected)", "money"],
                    ["principal_paid", "Repaid (verified)", "money"], ["principal_pending", "Pending transfer", "money"],
                    ["margin_payable", "Margin payable", "money"], ["due_date", "Due date", "text"], ["status", "Status", "text"]]
            summary = {"Principal due": inr(sum(r["principal_due"] for r in rows)), "Repaid": inr(sum(r["principal_paid"] for r in rows))}
    elif key in ("wheelspa_margin", "sales_collections", "gst"):
        sq = {**q, "status": {"$in": SOLD}}
        if f or t:
            sq["invoice_date"] = rng(f, t)
        rows = await _all("sales", sq, [("invoice_date", 1)])
        for r in rows:
            p = r["pricing"]
            r.update({"base": p["base"], "margin": p["margin"], "income_tax": p.get("income_tax", 0), "investor_margin": p["investor_margin"], "wheelspa_margin": p["wheelspa_margin"],
                      "taxable": p["taxable"], "gst": p["gst"], "gst_rate": f"{p['gst_bp']/100:g}%", "total": p["total"],
                      "outstanding": p["total"] - r["collected"]})
        if key == "wheelspa_margin":
            cols = [["order_no", "Order", "text"], ["invoice_date", "Invoice date", "text"], ["roll_serial", "Roll", "text"], ["product", "Product", "text"],
                    ["base", "Base", "money"], ["margin", "Margin 30%", "money"], ["income_tax", "Income tax", "money"],
                    ["wheelspa_margin", "Wheelspa share", "money"],
                    ["ws_recognized", "Recognised (collected)", "money"]]
            summary = {"Wheelspa margin": inr(sum(r["wheelspa_margin"] for r in rows)), "Recognised": inr(sum(r["ws_recognized"] for r in rows))}
        elif key == "gst":
            cols = [["invoice_no", "Invoice", "text"], ["invoice_date", "Date", "text"], ["customer_name", "Customer", "text"],
                    ["taxable", "Taxable value", "money"], ["gst_rate", "Rate", "text"], ["gst", "GST", "money"],
                    ["gst_collected", "GST collected", "money"], ["total", "Invoice total", "money"]]
            summary = {"Output GST": inr(sum(r["gst"] for r in rows)), "GST collected": inr(sum(r["gst_collected"] for r in rows))}
        else:
            cols = [["order_no", "Order", "text"], ["invoice_date", "Invoice date", "text"], ["customer_name", "Customer", "text"],
                    ["roll_serial", "Roll", "text"], ["total", "Invoice total", "money"], ["collected", "Collected", "money"],
                    ["outstanding", "Receivable", "money"], ["status", "Status", "text"]]
            summary = {"Invoiced": inr(sum(r["total"] for r in rows)), "Collected": inr(sum(r["collected"] for r in rows))}
    elif key == "inventory":
        rq = {**q, **({"purchase_date": rng(f, t)} if f or t else {})}
        rows = await _all("rolls", rq, [("roll_id", 1)])
        today = datetime.fromisoformat(today_ist())
        for r in rows:
            r["age_days"] = (today - datetime.fromisoformat(r["purchase_date"])).days
            r["investors"] = ", ".join(f"{x['investor_name']} {x['pct_bp']/100:g}%" for x in r["funding"])
        cols = [["roll_id", "Roll ID", "text"], ["po_number", "PO", "text"], ["brand", "Brand", "text"], ["series", "Series", "text"],
                ["purchase_date", "Purchased", "text"], ["status", "Status", "text"], ["age_days", "Age (days)", "int"],
                ["landed_cost", "Landed cost", "money"], ["base_price", "Agreed base", "money"], ["investors", "Funded by", "text"]]
        summary = {"Rolls": len(rows), "In stock": sum(r["status"] == "in_stock" for r in rows), "Sold": sum(r["status"] == "sold" for r in rows)}
    elif key == "purchases":
        rows = await _all("purchases", {**q, **({"purchase_date": rng(f, t)} if f or t else {})}, [("purchase_date", 1)])
        cols = [["po_number", "PO", "text"], ["purchase_date", "Date", "text"], ["supplier", "Supplier", "text"], ["supplier_invoice", "Supplier inv", "text"],
                ["qty", "Qty", "int"], ["subtotal", "Subtotal", "money"], ["gst", "GST", "money"], ["total", "Total", "money"],
                ["amount_paid", "Paid", "money"], ["payment_status", "Payment", "text"]]
        summary = {"Total procurement": inr(sum(r["total"] for r in rows)), "Paid": inr(sum(r["amount_paid"] for r in rows))}
    elif key == "outstanding_principal":
        iq = {} if can(m, "staff") else {"id": m.get("investor_id") or "__none__"}
        if investor_id and can(m, "staff"):
            iq = {"id": investor_id}
        rows = []
        for inv in await _all("investors", iq, [("code", 1)]):
            rows.append({"code": inv["code"], "name": inv["name"], **balances(inv)})
        cols = [["code", "Code", "text"], ["name", "Investor", "text"], ["invested", "Invested", "money"], ["allocated", "Allocated", "money"],
                ["recovered", "Recovered", "money"], ["repaid", "Repaid", "money"], ["outstanding", "Outstanding", "money"],
                ["available", "Available", "money"], ["deployed", "In inventory", "money"]]
    elif key == "monthly_reconciliation":
        iid = investor_id if can(m, "staff") else m.get("investor_id")
        if not iid:
            first = await db.investors.find_one({}, {"id": 1}, sort=[("code", 1)])
            iid = first["id"] if first else None
        if not iid:
            raise HTTPException(400, "Select an investor")
        cols, rows, summary = await monthly_reco(iid, month or today_ist()[:7])
    else:
        if not can(m, "audit"):
            raise HTTPException(403, "Not allowed")
        rows = await _all("audit", {"ts": rng(f, t)} if f or t else {}, [("ts", -1)])
        for r in rows:
            r["details"] = str(r.get("details") or "")[:160]
        cols = [["id", "ID", "text"], ["ts", "Timestamp (UTC)", "text"], ["user", "User", "text"], ["role", "Role", "text"],
                ["action", "Action", "text"], ["entity", "Entity", "text"], ["details", "Details", "text"]]
    return {"key": key, "title": REPORTS[key], "columns": cols, "rows": rows, "summary": summary,
            "period": f"{f or 'start'} to {t or 'today'}" if key != "monthly_reconciliation" else (month or today_ist()[:7]),
            "generated_at": datetime.now(IST).strftime("%d %b %Y, %I:%M %p IST")}


def sale_statement(sale, settlements, transfers):
    p = sale["pricing"]
    pr_rec = sum(s["principal_recoverable"] for s in settlements)
    pr_paid = sum(s["principal_paid"] for s in settlements)
    mg_paid = sum(s["margin_paid"] for s in settlements)
    verified = [t for t in transfers if t["status"] == "verified"]
    statuses = ", ".join(sorted({settlement_status(s, sale) for s in settlements})) or "-"
    items = [
        ("Sales order ID", sale["order_no"]), ("Customer", f"{sale['customer_name']} {sale.get('customer_code') or ''}".strip()),
        ("PPF roll ID", sale["roll_serial"]), ("Agreed base price", inr(p["base"])), ("Combined margin", inr(p["margin"])),
        (f"Income tax @ {p.get('tax_bp', 0)/100:g}% on margin", inr(p.get("income_tax", 0))), ("Net margin for distribution", inr(p.get("net_margin", p["margin"]))),
        ("Investor margin share", inr(p["investor_margin"])), ("Wheelspa margin share", inr(p["wheelspa_margin"])),
        ("Taxable selling value", inr(p["taxable"])), (f"GST @ {p['gst_bp']/100:g}%", inr(p["gst"])), ("Total invoice value", inr(p["total"])),
        ("Actual amount collected", inr(sale["collected"])), ("Principal recoverable (collected)", inr(pr_rec)),
        ("Principal repaid (verified)", inr(pr_paid)), ("Investor margin payable", inr(sum(s["margin_recoverable"] for s in settlements) - mg_paid)),
        ("Investor margin paid (verified)", inr(mg_paid)), ("Wheelspa margin payable/recognised", inr(sale["ws_recognized"])),
        ("Total transferred to investor", inr(pr_paid + mg_paid)),
        ("Bank transaction reference(s)", ", ".join(f"{t['kind']}: {t['bank_ref']}" for t in verified) or "-"),
        ("Settlement status", statuses),
        ("Outstanding after settlement", inr(sum(s["principal_due"] + s["margin_due"] - s["principal_paid"] - s["margin_paid"] for s in settlements))),
    ]
    return {"key": "statement", "title": f"Settlement Statement — {sale['order_no']}",
            "columns": [["k", "Particulars", "text"], ["v", "Value (INR)", "text"]],
            "rows": [{"k": k, "v": v} for k, v in items], "summary": {"Invoice": sale.get("invoice_no") or "not issued"},
            "period": sale["order_date"], "generated_at": datetime.now(IST).strftime("%d %b %Y, %I:%M %p IST")}


def _cell(r, c):
    v = r.get(c[0], "")
    if c[2] == "money":
        return inr(v)
    return "" if v is None else str(v)


def to_xlsx(rep) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws.append([rep["title"]])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([f"S & A Ventures · Period: {rep['period']} · Generated {rep['generated_at']}"])
    for k, v in rep["summary"].items():
        ws.append([k, str(v)])
    ws.append([])
    ws.append([c[1] + (" (₹)" if c[2] == "money" else "") for c in rep["columns"]])
    for cell in ws[ws.max_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="B45309")
    for r in rep["rows"]:
        ws.append([(r.get(c[0], 0) or 0) / 100 if c[2] == "money" else (r.get(c[0]) if c[2] == "int" else _cell(r, c)) for c in rep["columns"]])
        for i, c in enumerate(rep["columns"]):
            if c[2] == "money":
                ws.cell(row=ws.max_row, column=i + 1).number_format = "#,##0.00"
    for i in range(len(rep["columns"])):
        ws.column_dimensions[chr(65 + i)].width = 20
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(rep) -> bytes:
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24)
    ss = getSampleStyleSheet()
    fix = lambda s: str(s).replace("₹", "Rs. ").replace("−", "-").replace("—", "-").replace("·", "|")
    els = [Paragraph(fix(rep["title"]), ss["Title"]),
           Paragraph(fix(f"S & A Ventures | Period: {rep['period']} | Generated {rep['generated_at']}"), ss["Normal"]),
           Spacer(1, 6)]
    for k, v in rep["summary"].items():
        els.append(Paragraph(fix(f"<b>{k}:</b> {v}"), ss["Normal"]))
    els.append(Spacer(1, 8))
    cell = ss["BodyText"].clone("c", fontSize=7, leading=8.5)
    data = [[Paragraph(fix(c[1] + (" (Rs.)" if c[2] == "money" else "")), cell) for c in rep["columns"]]]
    data += [[Paragraph(fix(_cell(r, c)), cell) for c in rep["columns"]] for r in rep["rows"]] or [["No records"] + [""] * (len(rep["columns"]) - 1)]
    tbl = Table(data, repeatRows=1)
    tbl.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#FDE68A")),
                             ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#CBD5E1")),
                             ("VALIGN", (0, 0), (-1, -1), "TOP"),
                             ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F8FAFC")])]))
    els += [tbl, Spacer(1, 10), Paragraph("System-generated statement. GST shown separately and excluded from distributable margin. "
                                          "Accounting treatment subject to review by the company's Chartered Accountant.", ss["Italic"])]
    doc.build(els)
    return buf.getvalue()


# ---------- Dashboard & reconciliation ----------
async def dashboard(m):
    staff = can(m, "staff")
    q = scope(m)
    iq = {} if staff else {"id": m.get("investor_id") or "__none__"}
    invs = await _all("investors", iq)
    B = {k: sum(balances(i)[k] for i in invs) for k in balances({}).keys()}
    rolls = await _all("rolls", q)
    sales = await _all("sales", {**q, "status": {"$in": SOLD}})
    sts = [await enrich_settlement(s) for s in await _all("settlements", {**q, "active": True})]
    trs = await _all("transfers", {**q, "status": "initiated"})
    counts = {}
    for r in rolls:
        counts[r["status"]] = counts.get(r["status"], 0) + 1
    sc = {}
    for s in sts:
        sc[s["status"]] = sc.get(s["status"], 0) + 1
    overdue = sc.get("overdue", 0)
    pending = sum(v for k, v in sc.items() if k != "settled")
    today = datetime.fromisoformat(today_ist())
    aging = {"0-30": 0, "31-60": 0, "61-90": 0, "90+": 0}
    for r in rolls:
        if r["status"] == "in_stock":
            a = (today - datetime.fromisoformat(r["purchase_date"])).days
            aging["0-30" if a <= 30 else "31-60" if a <= 60 else "61-90" if a <= 90 else "90+"] += 1
    mine = m.get("investor_id")
    unsold_val = sum(f["principal"] for r in rolls if r["status"] in ("in_stock", "reserved") for f in r["funding"] if staff or f["investor_id"] == mine)
    stock_cost = sum(r["landed_cost"] for r in rolls if r["status"] in ("in_stock", "reserved"))
    months = {}
    for s in sales:
        k = (s.get("invoice_date") or s["order_date"])[:7]
        mo = months.setdefault(k, {"month": k, "base": 0, "investor_margin": 0, "wheelspa_margin": 0, "gst": 0, "collected": 0})
        for f in ("base", "investor_margin", "wheelspa_margin", "gst"):
            mo[f] += s["pricing"][f]
        mo["collected"] += s["collected"]
    products = {}
    for s in sales:
        pr = products.setdefault(s["product"] or "PPF", {"product": s["product"] or "PPF", "rolls": 0, "revenue": 0, "investor_margin": 0, "wheelspa_margin": 0})
        pr["rolls"] += 1
        pr["revenue"] += s["pricing"]["taxable"]
        pr["investor_margin"] += s["pricing"]["investor_margin"]
        pr["wheelspa_margin"] += s["pricing"]["wheelspa_margin"]
    total_sales = sum(s["pricing"]["total"] for s in sales)
    collected = sum(s["collected"] for s in sales)
    k = lambda key, label, value, kind="money", tone="neutral", sub="": {"key": key, "label": label, "value": value, "kind": kind, "tone": tone, "sub": sub}
    if staff:
        purchases = await _all("purchases", {})
        kpis = [k("funding", "Total investor funding", B["invested"]),
                k("procurement", "Procurement expenditure", sum(p["total"] for p in purchases)),
                k("available", "Available funds", B["available"], tone="amber" if B["available"] < 3125000 else "neutral"),
                k("inventory", "Inventory valuation (cost)", stock_cost, sub=f"Principal value ₹{inr(unsold_val)}"),
                k("sales", "Total sales (invoiced)", total_sales),
                k("collections", "Customer collections", collected, tone="green"),
                k("receivables", "Outstanding receivables", total_sales - collected, tone="amber" if total_sales > collected else "green"),
                k("principal_outstanding", "Investor principal outstanding", B["outstanding"]),
                k("margin_payable", "Investor margin payable", B["margin_payable"], tone="amber" if B["margin_payable"] else "green"),
                k("ws_margin", "Wheelspa margin earned", sum(s["ws_recognized"] for s in sales), tone="green"),
                k("pending_transfers", "Pending bank transfers", sum(t["amount"] for t in trs), tone="amber" if trs else "green", sub=f"{len(trs)} awaiting verification"),
                k("overdue", "Overdue settlements", overdue, kind="int", tone="red" if overdue else "green")]
    else:
        kpis = [k("invested", "Total principal invested", B["invested"]),
                k("utilised", "Funds utilised for purchases", B["allocated"]),
                k("recovered", "Principal recovered (collected)", B["recovered"], tone="green"),
                k("outstanding", "Principal outstanding", B["outstanding"], tone="amber" if B["outstanding"] else "green"),
                k("available", "Available unutilised funds", B["available"]),
                k("unsold", "Unsold inventory value", unsold_val),
                k("margin_earned", "Investor margin earned", B["margin_earned"], tone="green"),
                k("margin_paid", "Investor margin paid", B["margin_paid"], tone="green"),
                k("margin_payable", "Investor margin payable", B["margin_payable"], tone="amber" if B["margin_payable"] else "green"),
                k("transferred", "Total transferred to you", B["total_transferred"], tone="green"),
                k("pending_transfer", "Principal pending transfer", B["principal_pending_transfer"], tone="amber" if B["principal_pending_transfer"] else "green"),
                k("pending", "Pending settlements", pending, kind="int", tone="red" if overdue else ("amber" if pending else "green"), sub=f"{overdue} overdue")]
    recent = sorted(sts, key=lambda s: s.get("last_collection_date") or s["created_at"][:10], reverse=True)[:8]
    moves = await db.ledger.find({**q, "type": {"$in": ["CAPITAL_RECEIVED", "PRINCIPAL_REPAYMENT"]}}, {"_id": 0}).sort("txn_id", -1).to_list(12)
    names = {i["id"]: i["name"] for i in invs}
    for e in moves:
        e["investor_name"] = names.get(e["investor_id"], "Deleted investor")
    profit = {f: sum(s["pricing"].get(f, 0) for s in sales) for f in ("base", "margin", "income_tax", "net_margin", "investor_margin", "wheelspa_margin", "gst", "total")}
    profit["rolls"] = len(sales)
    aq = {"status": "approved", **({} if staff else {"investor_id": mine or "__none__"})}
    ag = await db.agreements.find_one(aq, {"_id": 0}, sort=[("effective_from", -1), ("version", -1)])
    example = None
    if ag:
        b = ag["default_base_price"]
        example = {**compute_pricing(b, [{"investor_id": ag["investor_id"], "principal": b}], {ag["investor_id"]: ag}),
                   "margin_bp": ag["margin_bp"], "investor_share_bp": ag["investor_share_bp"], "base_source": ag.get("base_source", "")}
    capital = {"added": B["invested"], "returned": B["repaid"], "outstanding": B["outstanding"]}
    low = {}
    for r in rolls if staff else []:
        b = low.setdefault(r["purchase_id"], {"purchase_id": r["purchase_id"], "po_number": r["po_number"], "batch_no": r["batch_no"],
                                              "product": f"{r['brand']} {r.get('series', '')}".strip(), "total": 0, "in_stock": 0})
        b["total"] += 1
        b["in_stock"] += r["status"] == "in_stock"
    low_stock = [b for b in low.values() if 1 <= b["in_stock"] <= 2]
    return {"kpis": kpis, "low_stock": low_stock, "inventory": {"purchased": len(rolls), "in_stock": counts.get("in_stock", 0), "reserved": counts.get("reserved", 0),
                                        "sold": counts.get("sold", 0), "exceptions": sum(counts.get(x, 0) for x in ("damaged", "returned", "disputed")),
                                        "purchase_cost": sum(r["landed_cost"] for r in rolls), "counts": counts},
            "aging": [{"bucket": a, "rolls": v} for a, v in aging.items()], "monthly": [months[x] for x in sorted(months)][-6:],
            "settlement_counts": sc, "products": list(products.values()), "recent": recent,
            "capital_moves": moves, "capital": capital, "profit": profit, "example": example}


async def reconcile(m, notify_on_fail=False):
    staff = can(m, "staff")
    iq = {} if staff else {"id": m.get("investor_id") or "__none__"}
    checks = []

    def chk(name, expected, actual, investor="All", kind="money"):
        checks.append({"investor": investor, "check": name, "expected": expected, "actual": actual, "ok": expected == actual, "kind": kind})

    for inv in await _all("investors", iq):
        b, L, nm = balances(inv), await ledger_sums(inv["id"]), inv["name"]
        sts = await _all("settlements", {"investor_id": inv["id"], "active": True})
        trs = await _all("transfers", {"investor_id": inv["id"], "status": "verified"})
        roll_alloc = sum(f["principal"] for r in await _all("rolls", {"investor_ids": inv["id"]}) for f in r["funding"] if f["investor_id"] == inv["id"])
        chk("Capital: balance vs ledger", L.get("CAPITAL_RECEIVED", 0), b["invested"], nm)
        chk("Allocation: balance vs roll-level funding", roll_alloc, b["allocated"], nm)
        chk("Allocation: ledger vs roll-level funding", roll_alloc, L.get("FUNDS_ALLOCATED", 0), nm)
        chk("Principal recovered: balance vs settlements", sum(s["principal_recoverable"] for s in sts), b["recovered"], nm)
        chk("Principal repaid: balance vs verified transfers", sum(t["amount"] for t in trs if t["kind"] == "principal"), b["repaid"], nm)
        chk("Principal repaid: ledger vs verified transfers", sum(t["amount"] for t in trs if t["kind"] == "principal"), L.get("PRINCIPAL_REPAYMENT", 0), nm)
        chk("Margin earned: balance vs settlements", sum(s["margin_recoverable"] for s in sts), b["margin_earned"], nm)
        chk("Margin paid: balance vs verified transfers", sum(t["amount"] for t in trs if t["kind"] == "margin"), b["margin_paid"], nm)
        chk("Capital = available + deployed + recovered", b["invested"], b["available"] + b["deployed"] + b["recovered"], nm)
        chk("Repayments do not exceed recovered principal", True, b["repaid"] <= b["recovered"], nm, "bool")
        chk("Available funds not negative", True, b["available"] >= 0, nm, "bool")
    q = scope(m)
    dup = await db.sales.aggregate([{"$match": {**q, "active": True}}, {"$group": {"_id": "$roll_ref", "n": {"$sum": 1}}}, {"$match": {"n": {"$gt": 1}}}]).to_list(100)
    chk("Rolls with more than one active sale", 0, len(dup), kind="int")
    sold = await db.rolls.count_documents({**q, "status": "sold"})
    chk("Sold rolls vs invoiced orders", await db.sales.count_documents({**q, "status": {"$in": SOLD}}), sold, kind="int")
    bad = [r["roll_id"] for r in await _all("rolls", q) if sum(f["principal"] for f in r["funding"]) != r["base_price"] or not r["funding"]]
    chk("Rolls with incomplete funding allocation", 0, len(bad), kind="int")
    nv = await db.transfers.count_documents({**q, "status": "verified", "$or": [{"bank_ref": ""}, {"doc_id": None}]})
    chk("Verified transfers missing bank proof", 0, nv, kind="int")
    ok = all(c["ok"] for c in checks)
    if not ok and notify_on_fail and staff:
        await notify("staff", "mismatch", "Reconciliation discrepancy / inventory mismatch",
                     f"{sum(not c['ok'] for c in checks)} check(s) failed", "/reconciliation")
    return {"ok": ok, "checks": checks, "run_at": now_iso()}
